"""Read back the entire native Google SBER pilot export against its frozen source."""
import argparse,collections,datetime as dt,gzip,hashlib,json,math
from pathlib import Path
import openpyxl
from birzha.application.data_standard import CANDLE_COLUMNS,LOGICAL_SHEETS,migrate_research_row

def literal(x):
 if x is None:return ''
 if isinstance(x,(dt.datetime,dt.date,dt.time)):return x.isoformat()
 if isinstance(x,float) and x.is_integer():return int(x)
 return x

def packed(x):return json.dumps(x,ensure_ascii=False,sort_keys=True,separators=(',',':'),allow_nan=False)
def digest_row(d,row):d.update((packed([literal(v) for v in row])+'\n').encode())
def main():
 p=argparse.ArgumentParser();p.add_argument('--source',type=Path,required=True);p.add_argument('--pilot',type=Path,required=True);p.add_argument('--audit-summary',type=Path,required=True);p.add_argument('--indicators',type=Path,required=True);p.add_argument('--manifest',type=Path,required=True);p.add_argument('--output',type=Path,required=True);a=p.parse_args()
 source=openpyxl.load_workbook(a.source,read_only=True,data_only=False);target=openpyxl.load_workbook(a.pilot,read_only=True,data_only=False);audit=json.loads(a.audit_summary.read_text(encoding='utf-8'));book=next(b for b in audit['books'] if b['symbol']=='SBER');manifest=json.loads(a.manifest.read_text(encoding='utf-8'));errors=[];checks=[];error_counts=collections.Counter()
 def error(kind,**fields):
  error_counts[kind]+=1
  if len(errors)<40:errors.append(dict(kind=kind,**fields))
 if target.sheetnames[:10]!=list(LOGICAL_SHEETS):error('LOGICAL_SHEET_ORDER')
 if not set(source.sheetnames).issubset(target.sheetnames):error('MISSING_OLD_SHEETS')
 unchanged=0;extended=0
 for name,info in book['sheets'].items():
  if name not in target.sheetnames:continue
  width=len(info['header']);count=info['actual_last_nonempty_row'];oldhash=hashlib.sha256();newhash=hashlib.sha256();sheetrows=0
  if not width:checks.append(dict(sheet=name,rows=0,preserved=True));continue
  before=list(source[name].iter_rows(min_row=1,max_row=count,max_col=width,values_only=True));after=list(target[name].iter_rows(min_row=1,max_row=count,max_col=40 if name in ('D1','H1','M15') else width,values_only=True))
  if name in ('D1','H1','M15') and list(after[0])!=CANDLE_COLUMNS:error('CANDLE_HEADER',sheet=name)
  for ordinal,(left,right) in enumerate(zip(before,after),1):
   old=[literal(v) for v in left];new=[literal(v) for v in right[:width]]
   digest_row(oldhash,old);digest_row(newhash,new)
   if old!=new:error('SOURCE_CELL_CHANGED',sheet=name,row=ordinal)
   if ordinal>1 and name in ('D1','H1','M15'):
    expected=migrate_research_row(info['header'],list(left),'SBER',name)
    if [literal(v) for v in expected]!=[literal(v) for v in right]:error('EXTENDED_METADATA_CHANGED',sheet=name,row=ordinal)
    extended+=1
   sheetrows+=1
  if oldhash.hexdigest()!=newhash.hexdigest():error('SOURCE_HASH_MISMATCH',sheet=name)
  checks.append(dict(sheet=name,rows=sheetrows,before_sha256=oldhash.hexdigest(),after_sha256=newhash.hexdigest(),preserved=oldhash.hexdigest()==newhash.hexdigest()));unchanged+=sheetrows
 expected={}
 for line in gzip.open(a.indicators,'rt',encoding='utf-8'):
  r=json.loads(line)
  if r['input_scope']!='GOOGLE_SNAPSHOT' or r['symbol']!='SBER':continue
  r['formula_id']=manifest['formula_id'];r['indicator_status']=packed(r['indicator_status']);expected[r['record_key']]=[r.get(k) for k in manifest['columns']]
 header=None;seen=set();numeric=0;rounded=0;max_abs=0.0;max_rel=0.0;expected_hash=hashlib.sha256();actual_hash=hashlib.sha256();metadata_expected=hashlib.sha256();metadata_actual=hashlib.sha256()
 for row_no,raw in enumerate(target['INDICATORS'].iter_rows(max_col=34,values_only=True),1):
  row=list(raw)
  if not any(v is not None and v!='' for v in row):continue
  if header is None:
   header=row
   if header!=manifest['columns']:error('INDICATOR_HEADER')
   continue
  key=row[0]
  if key in seen:error('DUPLICATE_INDICATOR_KEY',row=row_no)
  seen.add(key)
  if key not in expected:error('UNKNOWN_INDICATOR_KEY',row=row_no);continue
  want=expected[key];canon_expected=[];canon_actual=[]
  for col,(x,y) in enumerate(zip(want,row)):
   if isinstance(x,(float,int)) and not isinstance(x,bool):
    numeric+=1
    if isinstance(y,bool) or not isinstance(y,(int,float)) or not math.isfinite(y):error('NONNUMERIC_INDICATOR',row=row_no,column=col);continue
    difference=abs(x-y);max_abs=max(max_abs,difference);max_rel=max(max_rel,difference/max(1e-300,abs(x)))
    if x!=y:rounded+=1
    if not math.isclose(x,y,rel_tol=2e-14,abs_tol=1e-12):error('INDICATOR_VALUE_MISMATCH',row=row_no,column=col)
    canon_expected.append(float(format(x,'.11g')));canon_actual.append(float(format(y,'.11g')))
   else:
    if literal(x)!=literal(y):error('INDICATOR_METADATA_MISMATCH',row=row_no,column=col)
    canon_expected.append(literal(x));canon_actual.append(literal(y))
  digest_row(expected_hash,canon_expected);digest_row(actual_hash,canon_actual);digest_row(metadata_expected,want[:20]);digest_row(metadata_actual,row[:20])
 if len(seen)!=len(expected) or len(seen)!=43727:error('INDICATOR_COUNT',expected=len(expected),actual=len(seen))
 if metadata_expected.hexdigest()!=metadata_actual.hexdigest():error('INDICATOR_PROVENANCE_HASH')
 if expected_hash.hexdigest()!=actual_hash.hexdigest():error('NORMALIZED_INDICATOR_HASH')
 result=dict(status='PILOT_BACKREAD_PASS' if not errors else 'PILOT_BACKREAD_FAIL',scope='FULL_SBER_NATIVE_COPY_NOT_CANONICAL_BOOK',source_sha256=hashlib.sha256(a.source.read_bytes()).hexdigest(),pilot_sha256=hashlib.sha256(a.pilot.read_bytes()).hexdigest(),old_sheet_checks=checks,old_rows_including_headers_preserved=unchanged,extended_candle_rows_checked=extended,indicator_rows=len(seen),indicator_numeric_cells_checked=numeric,google_numeric_rounding_cells=rounded,max_abs_difference=max_abs,max_relative_difference=max_rel,numeric_tolerance=dict(relative=2e-14,absolute=1e-12),normalized_hash_significant_digits=11,indicator_expected_sha256=expected_hash.hexdigest(),indicator_actual_sha256=actual_hash.hexdigest(),provenance_expected_sha256=metadata_expected.hexdigest(),provenance_actual_sha256=metadata_actual.hexdigest(),old_ohlcv_changes=0 if not any(e['kind']=='SOURCE_CELL_CHANGED' for e in errors) else None,errors=errors,error_counts=dict(error_counts),strict_forecast_admission='NOT_GRANTED')
 source.close();target.close();a.output.write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8');print(json.dumps({k:result[k] for k in ('status','extended_candle_rows_checked','indicator_rows','google_numeric_rounding_cells','max_abs_difference','errors')},ensure_ascii=False));raise SystemExit(bool(errors))
if __name__=='__main__':main()
