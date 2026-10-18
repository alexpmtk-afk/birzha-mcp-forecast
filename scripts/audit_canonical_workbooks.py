"""Read-only canonical XLSX audit. Does not open HOME or edit source books."""
from __future__ import annotations
import argparse, collections, datetime as dt, gzip, hashlib, json, math
from pathlib import Path
import openpyxl
from openpyxl.utils.datetime import from_excel

SYMBOLS=("SBER","Si","BR","GOLD","IMOEX","RTSI")
MOSCOW=dt.timezone(dt.timedelta(hours=3))
def present(x): return x is not None and x!=""
def numeric(x): return isinstance(x,(int,float)) and not isinstance(x,bool) and math.isfinite(x)
def literal(x):
    if isinstance(x,(dt.datetime,dt.date,dt.time)): return x.isoformat()
    if isinstance(x,float) and not math.isfinite(x): return {"nonfinite":repr(x)}
    return x
def packed(x): return json.dumps(x,ensure_ascii=False,sort_keys=True,separators=(",",":"),allow_nan=False)
def timestamp(x):
    if isinstance(x,dt.datetime): p=x
    elif isinstance(x,dt.date): p=dt.datetime.combine(x,dt.time())
    elif isinstance(x,str):
        try: p=dt.datetime.fromisoformat(x.replace("Z","+00:00"))
        except ValueError: return None
    else: return None
    return p.replace(tzinfo=MOSCOW) if p.tzinfo is None else p
def row_date(x,fallback,epoch):
    p=timestamp(x)
    if not p and numeric(x) and 30000<x<80000: p=timestamp(from_excel(x,epoch))
    p=p or timestamp(fallback)
    return p.date().isoformat() if p else None

def candles(symbol,tf,header,rows,epoch,issue,staging):
    counts=collections.Counter(); keys={}; identities={}; intervals=collections.defaultdict(list); dates=[]
    fields={k:collections.Counter() for k in ("contract_code","secid","source_id","source_kind","research_status","available_at_confidence","source_interval","source_candles","is_complete")}
    missing=collections.Counter()
    for row_no,values in rows:
        record=dict(zip(header,values)); key=record.get("record_key")
        if not isinstance(key,str) or not key:
            counts["rows_without_record_key"]+=1
            issue(dict(symbol=symbol,sheet=tf,row=row_no,kind="MISSING_RECORD_KEY")); continue
        counts["records"]+=1; parts=key.split("|",3)
        secid=record.get("secid") or record.get("contract_code") or (parts[2] if len(parts)==4 else None)
        br=record.get("bar_start_time") or record.get("begin") or (parts[3] if len(parts)==4 else None)
        er=record.get("bar_end_time") or record.get("end"); begin,end=timestamp(br),timestamp(er)
        date=row_date(record.get("calendar_date"),br,epoch)
        if date: dates.append(date)
        raw_sha=hashlib.sha256(packed([literal(v) for v in values]).encode()).hexdigest()
        payload=tuple(packed(literal(record.get(k))) for k in ("open","high","low","close","volume","turnover"))
        if key in keys:
            counts["duplicate_record_keys"]+=1; conflict=keys[key][1]!=raw_sha; counts["different_payload_same_key"]+=int(conflict)
            issue(dict(symbol=symbol,sheet=tf,row=row_no,other_row=keys[key][0],key=key,kind="KEY_CONFLICT" if conflict else "IDENTICAL_DUPLICATE_KEY"))
        else: keys[key]=(row_no,raw_sha)
        if begin and secid:
            identity=(str(secid),tf,begin.isoformat())
            if identity in identities:
                old_row,old_payload=identities[identity]; conflict=payload!=old_payload
                counts["duplicate_semantic_identity"]+=1; counts["conflicting_semantic_prices_or_volume"]+=int(conflict)
                issue(dict(symbol=symbol,sheet=tf,row=row_no,other_row=old_row,identity=identity,kind="SEMANTIC_PRICE_CONFLICT" if conflict else "SEMANTIC_ALIAS_SAME_PRICES"))
            else: identities[identity]=(row_no,payload)
        else: counts["unparsed_identity"]+=1
        if begin and end:
            duration=(end-begin).total_seconds()+1; nominal={"D1":86400,"H1":3600,"M15":900}[tf]
            counts["nonpositive_intervals"]+=int(duration<=0); counts["shorter_than_nominal"]+=int(0<duration<nominal); counts["longer_than_nominal"]+=int(duration>nominal)
            intervals[str(secid)].append((begin,end,row_no))
        else: counts["unparsed_interval"]+=1
        o,h,l,c=(record.get(k) for k in ("open","high","low","close"))
        if not all(numeric(x) for x in (o,h,l,c)): counts["missing_or_invalid_ohlc"]+=1
        elif l>min(o,h,c) or h<max(o,l,c):
            counts["ohlc_invariant_violations"]+=1; issue(dict(symbol=symbol,sheet=tf,row=row_no,kind="OHLC_INVARIANT"))
        for field in ("volume","turnover","number_of_trades"):
            x=record.get(field); state="missing" if not present(x) else "invalid" if not numeric(x) else "negative" if x<0 else "zero" if x==0 else "positive"
            counts[field+"_"+state]+=1
        for field in ("available_at","observed_at","first_seen_at","trade_session_date","session_id","normalized_sha256","source_archive_sha256"):
            missing[field]+=int(not present(record.get(field)))
        for field,counter in fields.items(): counter[str(record.get(field)) if present(record.get(field)) else "<MISSING>"]+=1
        staging.write(packed(dict(symbol=symbol,timeframe=tf,source_row=row_no,source_record_sha256=raw_sha,source={k:literal(v) for k,v in record.items()},exact_secid=secid,parsed_begin=begin.isoformat() if begin else None,parsed_end=end.isoformat() if end else None))+"\n")
    for secid,bars in intervals.items():
        previous=None
        for begin,end,row_no in sorted(bars):
            if previous and begin<=previous[1]:
                counts["overlapping_intervals"]+=1; issue(dict(symbol=symbol,sheet=tf,row=row_no,other_row=previous[2],secid=secid,kind="OVERLAPPING_INTERVAL"))
            if previous is None or end>previous[1]: previous=(begin,end,row_no)
    return {**dict(counts),"unique_keys":len(keys),"unique_semantic_identities":len(identities),"first_date":min(dates) if dates else None,"last_date":max(dates) if dates else None,"observed_dates":len(set(dates)),"missing_fields":dict(missing),"absent_columns":[k for k in missing if k not in header],"categories":{k:dict(v) for k,v in fields.items()},"calendar_gaps":"NOT_ATTESTED_NO_INDEPENDENT_EXACT_SESSION_CALENDAR","m15_minute_replay":"NOT_ATTESTED_BY_XLSX_AUDIT"}

def audit_book(path,symbol,staging,additional,issue):
    book=openpyxl.load_workbook(path,read_only=True,data_only=False)
    result=dict(symbol=symbol,file=str(path.resolve()),file_sha256=hashlib.sha256(path.read_bytes()).hexdigest(),sheets={})
    for sheet in book.worksheets:
        header=None; rows=[]; actual_last=0; nonempty=0; digest=hashlib.sha256(); formulas=0
        for row_no,raw in enumerate(sheet.iter_rows(values_only=True),1):
            values=list(raw)
            while values and not present(values[-1]): values.pop()
            if not any(present(x) for x in values): continue
            nonempty+=1; actual_last=row_no
            digest.update((packed(dict(row=row_no,cells=[dict(type=type(v).__name__,value=literal(v)) for v in values]))+"\n").encode())
            formulas+=sum(isinstance(v,str) and v.startswith("=") for v in values)
            if header is None: header=[str(v) if v is not None else "" for v in values]
            else: rows.append((row_no,values))
        header=header or []; field_counts={}
        for i,name in enumerate(header):
            field_counts[name or f"<COL_{i+1}>"]=dict(nonblank=sum(i<len(r) and present(r[i]) for _,r in rows),numeric=sum(i<len(r) and numeric(r[i]) for _,r in rows))
        summary=dict(allocated_rows_not_coverage=sheet.max_row,allocated_columns=sheet.max_column,actual_nonempty_rows=nonempty,actual_data_rows=len(rows),actual_last_nonempty_row=actual_last,header=header,matrix_sha256=digest.hexdigest(),formula_cells=formulas,fields=field_counts)
        if sheet.title in ("D1","H1","M15"): summary["candles"]=candles(symbol,sheet.title,header,rows,book.epoch,issue,staging)
        else:
            for row_no,values in rows: additional.write(packed(dict(symbol=symbol,sheet=sheet.title,source_row=row_no,header=header,values=[literal(v) for v in values]))+"\n")
        result["sheets"][sheet.title]=summary
    book.close(); return result

def main():
    ap=argparse.ArgumentParser(); ap.add_argument("--snapshots",type=Path,required=True); ap.add_argument("--output",type=Path,required=True); args=ap.parse_args()
    args.output.mkdir(parents=True,exist_ok=True)
    result=dict(schema="GLOBAL_CANONICAL_XLSX_AUDIT_V1",scope="EXPORTED_GOOGLE_BOOKS_NOT_LIVE_DUCKDB",created_at_utc=dt.datetime.now(dt.timezone.utc).isoformat(),books=[])
    issues_count=0
    with gzip.open(args.output/"candle_records.jsonl.gz","wt",encoding="utf-8") as staging,gzip.open(args.output/"additional_records.jsonl.gz","wt",encoding="utf-8") as additional,gzip.open(args.output/"audit_anomalies.jsonl.gz","wt",encoding="utf-8") as issues:
        def issue(x):
            nonlocal issues_count
            issues_count+=1; issues.write(packed(x)+"\n")
        for symbol in SYMBOLS:
            record=audit_book(args.snapshots/(symbol+".xlsx"),symbol,staging,additional,issue); result["books"].append(record)
            print(symbol,{tf:record["sheets"][tf]["candles"]["records"] for tf in ("D1","H1","M15")},flush=True)
    result["anomaly_records"]=issues_count
    (args.output/"audit_summary.json").write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding="utf-8")
    print("AUDIT_COMPLETE anomalies="+str(issues_count),flush=True)
if __name__=="__main__": main()
